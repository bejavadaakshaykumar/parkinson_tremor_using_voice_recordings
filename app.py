import streamlit as st
import pandas as pd
import numpy as np
import plotly.express as px
import plotly.graph_objects as go
from plotly.subplots import make_subplots
from sklearn.preprocessing import MinMaxScaler
from sklearn.metrics import confusion_matrix, roc_curve, auc
from pathlib import Path
import hashlib
import json
from modeling import (
    MODEL_VERSION, _build_pipeline, prepare_data,
    train_model as train_grouped_model, predict_single as predict_raw,
)

import shap
import joblib
import warnings

import tempfile, os, random
import parselmouth
from datetime import datetime
from io import BytesIO
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

try:
    from audio_recorder_streamlit import audio_recorder
    RECORDER_AVAILABLE = True
except ImportError:
    RECORDER_AVAILABLE = False

try:
    from streamlit_shap import st_shap
    ST_SHAP_AVAILABLE = True
except ImportError:
    ST_SHAP_AVAILABLE = False

# ══════════════════════════════════════════════
#  CONSTANTS — Feature Engineering & Persistence
# ══════════════════════════════════════════════
APP_DIR = Path(__file__).resolve().parent
MODEL_PATH = APP_DIR / "parkinsons_pipeline_v7.pkl"
DATA_PATH = APP_DIR / "parkinsons_dataset.csv"
MODEL_CACHE_VERSION = MODEL_VERSION + hashlib.sha256(
    (APP_DIR / "modeling.py").read_bytes()
).hexdigest()[:12]

# ══════════════════════════════════════════════
#  PAGE CONFIG
# ══════════════════════════════════════════════
st.set_page_config(
    page_title="Parkinson's Voice Analytics",
    page_icon="🧠",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ══════════════════════════════════════════════
#  GLOBAL CSS
# ══════════════════════════════════════════════
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700;800;900&display=swap');

html, body, [data-testid="stAppViewContainer"],
[data-testid="stMain"], .main { background:#ffffff !important; }
[data-testid="stHeader"] { background:transparent !important; }
* { font-family:'Inter',sans-serif !important; }

/* ── Sidebar ── */
[data-testid="stSidebar"] {
    background: linear-gradient(175deg,#0d1b2a 0%,#1b2d4f 55%,#162340 100%) !important;
}
[data-testid="stSidebar"] * { color:#cbd5e1 !important; }
[data-testid="stSidebar"] h1,[data-testid="stSidebar"] h2,
[data-testid="stSidebar"] h3 { color:#f1f5f9 !important; }
[data-testid="stSidebar"] .stSelectbox label,
[data-testid="stSidebar"] .stSlider label,
[data-testid="stSidebar"] .stMultiSelect label {
    color:#94a3b8 !important; font-size:11px;
    letter-spacing:.08em; text-transform:uppercase; font-weight:600;
}

/* ── HERO ── */
.hero {
    background: linear-gradient(135deg,#0d1b2a 0%,#1e3a8a 45%,#7c3aed 100%);
    border-radius:24px; padding:48px 56px; margin-bottom:32px;
    position:relative; overflow:hidden;
    animation: fadeSlideDown .8s cubic-bezier(.22,1,.36,1) both;
    box-shadow: 0 24px 64px rgba(30,58,138,.25);
}
.hero::before {
    content:""; position:absolute; inset:0;
    background:radial-gradient(ellipse at 80% 20%,rgba(124,58,237,.35) 0%,transparent 60%);
}
.hero::after {
    content:""; position:absolute; right:-60px; top:-60px;
    width:260px; height:260px; border-radius:50%;
    background:rgba(255,255,255,.04);
}
.hero-emoji { font-size:3rem; display:block; margin-bottom:10px; animation:pulse 2.5s ease infinite; }
.hero-title { color:#fff; font-size:2.5rem; font-weight:900; margin:0 0 8px; letter-spacing:-.03em; }
.hero-sub   { color:#bfdbfe; font-size:1.05rem; margin:0; font-weight:400; }
.hero-badge {
    display:inline-block; margin-top:16px;
    background:rgba(255,255,255,.12); backdrop-filter:blur(8px);
    border:1px solid rgba(255,255,255,.2); border-radius:999px;
    padding:6px 18px; color:#e0f2fe; font-size:.8rem; font-weight:600;
    letter-spacing:.06em; text-transform:uppercase;
}

/* ── Pulse bar ── */
.pulse-bar {
    height:3px; border-radius:999px;
    background:linear-gradient(90deg,#6366f1,#8b5cf6,#06b6d4,#6366f1);
    background-size:300% 100%;
    animation:pulseSlide 3.5s linear infinite;
    margin:0 0 28px;
}
@keyframes pulseSlide { to { background-position:-300% 0; } }

/* ── KPI grid ── */
.kpi-grid { display:grid; grid-template-columns:repeat(4,1fr); gap:18px; margin-bottom:28px; }
.kpi-card {
    background:#fff; border:1px solid #e2e8f0; border-radius:18px;
    padding:24px 22px; box-shadow:0 4px 24px rgba(0,0,0,.06);
    animation:fadeUp .6s cubic-bezier(.22,1,.36,1) both;
    transition:transform .3s,box-shadow .3s; cursor:default;
    position:relative; overflow:hidden;
}
.kpi-card::before {
    content:""; position:absolute; top:0; left:0; right:0; height:3px;
    background:var(--accent,#6366f1);
    transform:scaleX(0); transform-origin:left;
    transition:transform .4s ease;
}
.kpi-card:hover::before { transform:scaleX(1); }
.kpi-card:hover { transform:translateY(-6px); box-shadow:0 16px 40px rgba(0,0,0,.12); }
.kpi-icon  { font-size:2rem; margin-bottom:10px; display:block; }
.kpi-label { font-size:10.5px; font-weight:700; letter-spacing:.1em;
             text-transform:uppercase; color:#94a3b8; margin-bottom:6px; }
.kpi-value { font-size:2.1rem; font-weight:900; color:#0f172a; line-height:1; }
.kpi-sub   { font-size:12px; color:#64748b; margin-top:6px; }

/* ── Section header ── */
.sec-hdr {
    display:flex; align-items:center; gap:10px;
    font-size:1.1rem; font-weight:800; color:#0f172a;
    margin:28px 0 14px; padding-bottom:12px;
    border-bottom:2px solid #ede9fe;
}

/* ── Chart card ── */
.chart-card {
    background:#fff; border:1px solid #e2e8f0; border-radius:18px;
    padding:12px 16px 16px; box-shadow:0 4px 20px rgba(0,0,0,.05);
    animation:fadeUp .7s cubic-bezier(.22,1,.36,1) both;
    margin-bottom:18px; transition:box-shadow .3s;
}
.chart-card:hover { box-shadow:0 10px 36px rgba(99,102,241,.1); }

/* ── Prediction box ── */
.pred-box {
    border-radius:18px; padding:28px 32px; text-align:center;
    animation:zoomIn .5s cubic-bezier(.22,1,.36,1) both;
    margin-bottom:16px;
}
.pred-pos { background:linear-gradient(135deg,#fef2f2,#ffe4e6); border:2px solid #fca5a5; }
.pred-neg { background:linear-gradient(135deg,#f0fdf4,#dcfce7); border:2px solid #86efac; }
.pred-title { font-size:1.6rem; font-weight:900; margin-bottom:8px; }
.pred-conf  { font-size:1rem; font-weight:500; color:#475569; }

/* ── Animations ── */
@keyframes fadeSlideDown {
    from { opacity:0; transform:translateY(-28px); }
    to   { opacity:1; transform:translateY(0); }
}
@keyframes fadeUp {
    from { opacity:0; transform:translateY(22px); }
    to   { opacity:1; transform:translateY(0); }
}
@keyframes zoomIn {
    from { opacity:0; transform:scale(.9); }
    to   { opacity:1; transform:scale(1); }
}
@keyframes pulse {
    0%,100% { transform:scale(1); }
    50%      { transform:scale(1.08); }
}

@media (max-width: 760px) {
    .hero { padding:28px 24px; border-radius:18px; }
    .hero-title { font-size:1.8rem; }
    .kpi-grid { grid-template-columns:repeat(2,1fr); gap:12px; }
    .kpi-card { padding:18px 14px; }
}
@media (prefers-reduced-motion: reduce) {
    *, *::before, *::after { animation:none !important; transition:none !important; }
}
/* ── Tabs ── */
[data-testid="stTabs"] [role="tab"] {
    font-weight:600; font-size:.88rem; color:#64748b;
    border-radius:10px 10px 0 0;
}
[data-testid="stTabs"] [role="tab"][aria-selected="true"] {
    color:#7c3aed; border-bottom:3px solid #7c3aed;
}
::-webkit-scrollbar { width:6px; }
::-webkit-scrollbar-track { background:#f8fafc; }
::-webkit-scrollbar-thumb { background:#cbd5e1; border-radius:999px; }
</style>
""", unsafe_allow_html=True)


# ══════════════════════════════════════════════
#  DATA & MODEL
# ══════════════════════════════════════════════
@st.cache_data(max_entries=4)
def load_data(csv_bytes):
    data, _, _, _, _ = prepare_data(pd.read_csv(BytesIO(csv_bytes)))
    return data


@st.cache_resource(show_spinner=False, max_entries=2)
def train_model(df, version=MODEL_CACHE_VERSION):
    """Cache the dataset-specific, nested patient-group evaluation and model."""
    progress = st.empty()
    try:
        return train_grouped_model(
            df, version=version, cache_path=MODEL_PATH, progress=progress.info,
        )
    finally:
        progress.empty()


# Load from the app directory, regardless of the shell's working directory.
if DATA_PATH.exists():
    csv_bytes = DATA_PATH.read_bytes()
else:
    st.title("Parkinson's Voice Analytics")
    st.info("Add parkinsons_dataset.csv beside app.py, or upload your training CSV below.")
    uploaded_data = st.file_uploader("Training dataset", type=["csv", "data"])
    if uploaded_data is None:
        st.stop()
    csv_bytes = uploaded_data.getvalue()

try:
    df = load_data(csv_bytes)
    required_acoustic = {
        "MDVP:Fo(Hz)", "MDVP:Fhi(Hz)", "MDVP:Flo(Hz)", "MDVP:Jitter(%)",
        "MDVP:Jitter(Abs)", "MDVP:RAP", "MDVP:PPQ", "Jitter:DDP", "MDVP:Shimmer",
        "MDVP:Shimmer(dB)", "Shimmer:APQ3", "Shimmer:APQ5", "MDVP:APQ", "Shimmer:DDA",
        "NHR", "HNR", "RPDE", "DFA", "spread1", "spread2", "D2", "PPE",
    }
    missing_acoustic = required_acoustic - set(df.columns)
    if missing_acoustic:
        raise ValueError("This dashboard needs the UCI acoustic columns: " + ", ".join(sorted(missing_acoustic)))
    with st.spinner("Loading or training the model. First-run nested validation can take several minutes…"):
        (
            pipeline, fitted_scaler, features, selected_features,
            cv_metrics, oof_preds, oof_probs, y_all, fitted_clf, hgb_model,
        ) = train_model(df, version=MODEL_CACHE_VERSION)
except (ValueError, OSError) as exc:
    st.error(f"Unable to train: {exc}")
    st.stop()

decision_threshold = pipeline.decision_threshold_
JITTER_COLS = ["MDVP:Jitter(%)", "MDVP:Jitter(Abs)", "MDVP:RAP", "MDVP:PPQ", "Jitter:DDP"]
SHIMMER_COLS = ["MDVP:Shimmer", "MDVP:Shimmer(dB)", "Shimmer:APQ3", "Shimmer:APQ5", "MDVP:APQ", "Shimmer:DDA"]
FREQ_COLS = ["MDVP:Fo(Hz)", "MDVP:Fhi(Hz)", "MDVP:Flo(Hz)"]
NONLINEAR_COLS = ["RPDE", "DFA", "spread1", "spread2", "D2", "PPE"]
RATIO_COLS = ["NHR", "HNR"]

PL = dict(
    paper_bgcolor="rgba(0,0,0,0)",
    plot_bgcolor="rgba(0,0,0,0)",
    font=dict(family="Inter,sans-serif", color="#1e293b"),
    margin=dict(l=8, r=8, t=44, b=8),
)
PAL = {"Parkinson's": "#7c3aed", "Healthy": "#10b981"}


def hex_to_rgba(hex_color, alpha=0.2):
    """Convert hex color to rgba string for Plotly compatibility."""
    h = hex_color.lstrip("#")
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    return f"rgba({r},{g},{b},{alpha})"


def predict_single(x_raw):
    """Apply all fitted transforms and the saved production threshold."""
    return predict_raw(pipeline, x_raw)


# ══════════════════════════════════════════════
#  BULLETPROOF SHAP HELPERS
# ══════════════════════════════════════════════
@st.cache_data(max_entries=2, show_spinner=False)
def global_shap_values(model_fingerprint, X, _model):
    transformed = _model.transform_for_explanation(X)
    return shap.TreeExplainer(_model.estimators_[0])(transformed).values


def _safe_shap_single(sv):
    """Extract SHAP values for a single sample's positive class.
    Handles 0D, 1D, and 2D arrays safely.
    """
    if sv.ndim == 2:
        return sv[:, 1]          # (features, 2) → positive class column
    elif sv.ndim == 0:
        return np.array([float(sv)])
    return sv                    # 1D → already correct


def _safe_shap_global(sv):
    """Extract global SHAP values for positive class.
    Handles 2D and 3D arrays safely.
    """
    if sv.ndim == 3:
        return sv[:, :, 1]       # (samples, features, 2) → positive class
    elif sv.ndim == 1:
        return sv.reshape(1, -1)
    return sv                    # 2D → already correct


def _shap_bar_fallback(sv, feature_names_list):
    """Render SHAP values as a Plotly horizontal bar chart (fallback)."""
    shap_df = pd.DataFrame({
        "Feature": feature_names_list,
        "SHAP Value": sv,
    }).sort_values("SHAP Value", key=abs, ascending=True)
    fig = px.bar(
        shap_df.tail(min(15, len(feature_names_list))),
        y="Feature", x="SHAP Value",
        orientation="h",
        title="SHAP Feature Contributions",
        color="SHAP Value",
        color_continuous_scale="RdBu_r",
        color_continuous_midpoint=0,
    )
    fig.update_layout(**PL, coloraxis_showscale=False, height=420)
    st.plotly_chart(fig, width="stretch")


def render_shap_plot(x_input, feature_names_list):
    """Render a SHAP waterfall/force plot for the dominant HGB base estimator.
    Handles 1D/2D/3D SHAP arrays safely for any shap library version."""
    try:
        explainer = shap.TreeExplainer(hgb_model)
        shap_values = explainer(x_input)
        sv = _safe_shap_single(shap_values[0].values)

        if ST_SHAP_AVAILABLE:
            try:
                base = shap_values[0].base_values
                if hasattr(base, '__len__') and len(base) > 1:
                    base = float(base[1])
                exp = shap.Explanation(
                    values=sv,
                    base_values=base,
                    data=shap_values[0].data,
                    feature_names=feature_names_list,
                )
                st_shap(shap.plots.force(exp), height=160)
            except Exception:
                _shap_bar_fallback(sv, feature_names_list)
        else:
            _shap_bar_fallback(sv, feature_names_list)
    except Exception as e:
        st.warning(f"SHAP explanation unavailable: {e}")


# ══════════════════════════════════════════════
#  SIDEBAR
# ══════════════════════════════════════════════
with st.sidebar:
    st.markdown("## 🧠 Parkinson's Analytics")
    st.markdown("---")
    st.markdown("#### 🎛️ Dashboard Filters")
    status_filter = st.multiselect(
        "Status", ["Parkinson's", "Healthy"], default=["Parkinson's", "Healthy"]
    )
    patient_filter = st.multiselect(
        "Patient ID", sorted(df["patient_id"].unique()), default=[]
    )
    st.markdown("---")
    st.markdown("#### 🔬 Feature Group")
    feat_group = st.selectbox(
        "Explore group",
        [
            "Jitter (Frequency Variation)",
            "Shimmer (Amplitude Variation)",
            "Fundamental Frequency",
            "Nonlinear Dynamics",
            "Noise Ratios",
        ],
    )
    st.markdown("---")
    st.caption(
        f"{len(df)} recordings · {df.patient_id.nunique()} patients · {len(selected_features)} selected features"
    )

mask = df["label"].isin(status_filter)
if patient_filter:
    mask &= df["patient_id"].isin(patient_filter)
dff = df[mask].copy()
if dff.empty:
    st.info("No recordings match these filters. Select a status or clear the patient filter.")
    st.stop()

# ══════════════════════════════════════════════
#  HERO + PULSE BAR
# ══════════════════════════════════════════════
st.markdown(
    f"""
<div class="hero">
  <span class="hero-emoji">🧠</span>
  <div class="hero-title">Parkinson's Voice Analytics</div>
  <div class="hero-sub">Explore acoustic patterns · {len(selected_features)} selected features · Patient-group validation</div>
  <span class="hero-badge">📊 {len(df)} Recordings · {df.patient_id.nunique()} Patients · Research Dashboard</span>
</div>
""",
    unsafe_allow_html=True,
)
st.markdown('<div class="pulse-bar"></div>', unsafe_allow_html=True)

# ══════════════════════════════════════════════
#  KPI CARDS
# ══════════════════════════════════════════════
total = len(dff)
pk_cnt = int((dff["status"] == 1).sum())
hl_cnt = int((dff["status"] == 0).sum())
pk_pct = round(pk_cnt / total * 100, 1) if total else 0
avg_hnr = round(float(dff["HNR"].mean()), 2) if total else 0

st.markdown(
    f"""
<div class="kpi-grid">
  <div class="kpi-card" style="--accent:#6366f1">
    <span class="kpi-icon">👥</span>
    <div class="kpi-label">Total Recordings</div>
    <div class="kpi-value">{total}</div>
    <div class="kpi-sub">voice samples in view</div>
  </div>
  <div class="kpi-card" style="--accent:#dc2626">
    <span class="kpi-icon">🔴</span>
    <div class="kpi-label">Parkinson's</div>
    <div class="kpi-value">{pk_cnt}</div>
    <div class="kpi-sub">{pk_pct}% of dataset</div>
  </div>
  <div class="kpi-card" style="--accent:#10b981">
    <span class="kpi-icon">🟢</span>
    <div class="kpi-label">Healthy</div>
    <div class="kpi-value">{hl_cnt}</div>
    <div class="kpi-sub">{round(100 - pk_pct, 1)}% of dataset</div>
  </div>
  <div class="kpi-card" style="--accent:#f59e0b">
    <span class="kpi-icon">🎙️</span>
    <div class="kpi-label">Avg HNR</div>
    <div class="kpi-value">{avg_hnr}</div>
    <div class="kpi-sub">Harmonics-to-Noise Ratio</div>
  </div>
</div>
""",
    unsafe_allow_html=True,
)

# ══════════════════════════════════════════════
#  TABS
# ══════════════════════════════════════════════
tab1, tab2, tab3, tab4, tab5, tab6 = st.tabs(
    [
        "📊 Overview",
        "🔬 Feature Deep Dive",
        "🤖 ML Model",
        "🧬 Predict",
        "📋 Data Explorer",
        "🎤 Voice Predict",
    ]
)

# ── TAB 1: OVERVIEW ──────────────────────────
with tab1:
    c1, c2 = st.columns(2)

    with c1:
        st.markdown('<div class="chart-card">', unsafe_allow_html=True)
        pie = dff["label"].value_counts().reset_index()
        pie.columns = ["Status", "Count"]
        fig = px.pie(
            pie,
            names="Status",
            values="Count",
            title="Class Distribution",
            color="Status",
            color_discrete_map=PAL,
            hole=0.52,
        )
        fig.update_traces(textinfo="percent+label", pull=[0.04, 0])
        fig.update_layout(**PL, legend=dict(orientation="h", y=-0.1))
        st.plotly_chart(fig, width="stretch")
        st.markdown("</div>", unsafe_allow_html=True)

    with c2:
        st.markdown('<div class="chart-card">', unsafe_allow_html=True)
        pt = (
            dff.groupby(["patient_id", "label"])["status"]
            .count()
            .reset_index()
        )
        pt.columns = ["Patient", "Status", "Recordings"]
        fig2 = px.bar(
            pt,
            x="Patient",
            y="Recordings",
            color="Status",
            title="Recordings per Patient",
            color_discrete_map=PAL,
            barmode="stack",
        )
        fig2.update_layout(**PL, xaxis_tickangle=45, xaxis_title="")
        st.plotly_chart(fig2, width="stretch")
        st.markdown("</div>", unsafe_allow_html=True)

    # Radar chart — core features only
    st.markdown('<div class="chart-card">', unsafe_allow_html=True)
    key_feats = [f for f in ["MDVP:Fo(Hz)", "HNR", "RPDE", "DFA", "PPE",
                              "spread1", "spread2", "D2"] if f in dff.columns]
    h_vals = dff[dff["status"] == 0][key_feats].mean().tolist()
    p_vals = dff[dff["status"] == 1][key_feats].mean().tolist()
    mms = MinMaxScaler()
    mat = mms.fit_transform(np.array([h_vals, p_vals]).T).T
    h_n, p_n = mat[0].tolist() + [mat[0][0]], mat[1].tolist() + [mat[1][0]]
    cats = key_feats + [key_feats[0]]
    fig3 = go.Figure()
    fig3.add_trace(
        go.Scatterpolar(
            r=h_n,
            theta=cats,
            fill="toself",
            name="Healthy",
            line_color="#10b981",
            fillcolor="rgba(16,185,129,.15)",
        )
    )
    fig3.add_trace(
        go.Scatterpolar(
            r=p_n,
            theta=cats,
            fill="toself",
            name="Parkinson's",
            line_color="#7c3aed",
            fillcolor="rgba(124,58,237,.15)",
        )
    )
    fig3.update_layout(
        **PL,
        title="Feature Radar — Normalised Mean Comparison",
        polar=dict(radialaxis=dict(visible=True, range=[0, 1])),
        legend=dict(orientation="h", y=-0.12),
    )
    st.plotly_chart(fig3, width="stretch")
    st.markdown("</div>", unsafe_allow_html=True)

    # Correlation heatmap — core features only
    st.markdown('<div class="chart-card">', unsafe_allow_html=True)
    corr_cols = [c for c in features if c in dff.columns]
    corr = dff[corr_cols + ["status"]].corr()
    fig4 = px.imshow(
        corr,
        text_auto=".2f",
        color_continuous_scale="RdBu_r",
        title="Feature Correlation Matrix (Core Features)",
        aspect="auto",
        zmin=-1,
        zmax=1,
    )
    fig4.update_layout(**PL, height=520)
    st.plotly_chart(fig4, width="stretch")
    st.markdown("</div>", unsafe_allow_html=True)

# ── TAB 2: FEATURE DEEP DIVE ─────────────────
with tab2:
    group_map = {
        "Jitter (Frequency Variation)": JITTER_COLS,
        "Shimmer (Amplitude Variation)": SHIMMER_COLS,
        "Fundamental Frequency": FREQ_COLS,
        "Nonlinear Dynamics": NONLINEAR_COLS,
        "Noise Ratios": RATIO_COLS,
    }
    cols_sel = group_map[feat_group]

    st.markdown('<div class="chart-card">', unsafe_allow_html=True)
    n = len(cols_sel)
    ncols = min(3, n)
    nrows = (n + ncols - 1) // ncols
    fig5 = make_subplots(
        rows=nrows, cols=ncols, subplot_titles=cols_sel, vertical_spacing=0.12
    )
    for idx, col in enumerate(cols_sel):
        r, c = divmod(idx, ncols)
        for lbl, color in PAL.items():
            sub = dff[dff["label"] == lbl]
            fig5.add_trace(
                go.Violin(
                    y=sub[col],
                    name=lbl,
                    fillcolor=hex_to_rgba(color, 0.2),
                    line_color=color,
                    box_visible=True,
                    meanline_visible=True,
                    showlegend=(idx == 0),
                ),
                row=r + 1,
                col=c + 1,
            )
    fig5.update_layout(
        **PL,
        title=f"{feat_group} — Distributions by Status",
        height=380 * nrows,
        violingap=0.3,
    )
    st.plotly_chart(fig5, width="stretch")
    st.markdown("</div>", unsafe_allow_html=True)

    st.markdown('<div class="chart-card">', unsafe_allow_html=True)
    melt = dff[cols_sel + ["label"]].melt(
        id_vars="label", var_name="Feature", value_name="Value"
    )
    fig6 = px.box(
        melt,
        x="Feature",
        y="Value",
        color="label",
        title=f"{feat_group} — Grouped Box Plot",
        color_discrete_map=PAL,
        points="outliers",
    )
    fig6.update_layout(**PL)
    st.plotly_chart(fig6, width="stretch")
    st.markdown("</div>", unsafe_allow_html=True)

    if len(cols_sel) >= 2:
        st.markdown('<div class="chart-card">', unsafe_allow_html=True)
        pair_x = st.selectbox("X axis", cols_sel, index=0)
        pair_y = st.selectbox("Y axis", cols_sel, index=min(1, len(cols_sel) - 1))
        fig7 = px.scatter(
            dff,
            x=pair_x,
            y=pair_y,
            color="label",
            marginal_x="histogram",
            marginal_y="histogram",
            color_discrete_map=PAL,
            opacity=0.75,
            title=f"{pair_x} vs {pair_y}",
        )
        fig7.update_layout(**PL)
        st.plotly_chart(fig7, width="stretch")
        st.markdown("</div>", unsafe_allow_html=True)

# ── TAB 3: ML MODEL ──────────────────────────
with tab3:
    st.markdown(
        '<div class="sec-hdr">🤖 Stacking Ensemble — Subject-Wise StratifiedGroupKFold CV</div>',
        unsafe_allow_html=True,
    )
    st.markdown(
        "Each base learner: **StandardScaler → SelectKBest (MI, 10) → SMOTEENN**. "
        "HGB + RF + SVC feed a cost-sensitive logistic stack. "
        "All evaluation, threshold-tuning, and stacking splits keep patients separate."
    )

    st.caption(
        "Metrics below are mean scores across five untouched outer folds. "
        "Thresholds are chosen with Youden's J on inner validation predictions only. "
        "Repeated recordings are grouped by patient; these are recording-level metrics."
    )
    st.metric("Saved decision threshold", f"{decision_threshold:.3f}")
    st.caption("The final model uses the mean of the five training-only thresholds. Scores are not calibrated disease probabilities.")

    # ── 4 KPI Cards: Accuracy, F1, Sensitivity, Specificity ──
    c1, c2, c3, c4 = st.columns(4)
    metrics_display = [
        ("Accuracy", f"{cv_metrics['accuracy']:.2f}%", "#6366f1", "🎯"),
        ("F1-Score", f"{cv_metrics['f1']:.2f}%", "#7c3aed", "⚡"),
        ("Sensitivity", f"{cv_metrics['sensitivity']:.2f}%", "#f59e0b", "🔬"),
        ("Specificity", f"{cv_metrics['specificity']:.2f}%", "#10b981", "🛡️"),
    ]
    for col, (label, val, color, icon) in zip([c1, c2, c3, c4], metrics_display):
        with col:
            st.markdown(
                f"""
            <div class="kpi-card" style="--accent:{color}">
              <span class="kpi-icon">{icon}</span>
              <div class="kpi-label">{label}</div>
              <div class="kpi-value" style="color:{color}">{val}</div>
              <div class="kpi-sub">Subject-wise 5-Fold CV</div>
            </div>""",
                unsafe_allow_html=True,
            )

    metric_names = ["accuracy", "f1", "sensitivity", "specificity"]
    with st.expander("Compare tuned decisions with the default 0.5 threshold", expanded=True):
        comparison = pd.DataFrame({
            "Metric": ["Accuracy", "F1-Score", "Sensitivity", "Specificity"],
            "Default 0.5 (%)": [cv_metrics["default_threshold_metrics"][m] for m in metric_names],
            "Tuned (%)": [cv_metrics[m] for m in metric_names],
        })
        st.dataframe(comparison.round(2), width="stretch", hide_index=True)
        st.caption("Same held-out scores, different decision thresholds. This comparison is not the old model baseline.")
    if all(cv_metrics[m] > 95 for m in metric_names):
        st.success("All four fold-mean metrics exceed 95% on this evaluation. Independent validation is still needed.")
    else:
        st.info("The >95% target is not met on all four metrics. The held-out results below show the measured performance.")

    # ── Per-fold performance table ──
    st.markdown('<div class="chart-card">', unsafe_allow_html=True)
    fold_df = pd.DataFrame({
        "Fold": [f"Fold {i+1}" for i in range(len(cv_metrics['fold_acc']))],
        "Accuracy (%)": [round(v * 100, 2) for v in cv_metrics['fold_acc']],
        "F1-Score (%)": [round(v * 100, 2) for v in cv_metrics['fold_f1']],
        "Sensitivity (%)": [round(v * 100, 2) for v in cv_metrics['fold_sens']],
        "Specificity (%)": [round(v * 100, 2) for v in cv_metrics['fold_spec']],
    })
    fold_df["Threshold"] = cv_metrics["fold_thresholds"]
    fold_df["Healthy recordings"] = [a["healthy_recordings"] for a in cv_metrics["fold_audit"]]
    fold_df["Parkinson's recordings"] = [a["parkinsons_recordings"] for a in cv_metrics["fold_audit"]]
    st.dataframe(fold_df.round(3), width="stretch", hide_index=True)
    st.download_button("Download validation audit", json.dumps(cv_metrics, indent=2),
                       file_name="validation_audit.json", mime="application/json")
    fig_fold = go.Figure()
    for metric_col, color in [
        ("Accuracy (%)", "#6366f1"),
        ("F1-Score (%)", "#7c3aed"),
        ("Sensitivity (%)", "#f59e0b"),
        ("Specificity (%)", "#10b981"),
    ]:
        fig_fold.add_trace(go.Bar(
            x=fold_df["Fold"], y=fold_df[metric_col],
            name=metric_col.replace(" (%)", ""),
            marker_color=color,
        ))
    fig_fold.update_layout(
        **PL,
        title="Per-Fold Performance (Subject-Wise StratifiedGroupKFold)",
        barmode="group",
        yaxis_title="Score (%)",
        yaxis_range=[0, 105],
        height=360,
        legend=dict(orientation="h", y=-0.15),
    )
    st.plotly_chart(fig_fold, width="stretch")
    st.markdown("</div>", unsafe_allow_html=True)

    # ── Confusion Matrix & ROC Curve ──
    c1, c2 = st.columns(2)
    with c1:
        st.markdown('<div class="chart-card">', unsafe_allow_html=True)
        cm = confusion_matrix(y_all, oof_preds)
        fig8 = px.imshow(
            cm,
            text_auto=True,
            x=["Pred Healthy", "Pred Parkinson's"],
            y=["Act Healthy", "Act Parkinson's"],
            color_continuous_scale="Purples",
            title="Confusion Matrix (Aggregated OOF)",
        )
        fig8.update_layout(**PL, height=340)
        st.plotly_chart(fig8, width="stretch")
        st.markdown("</div>", unsafe_allow_html=True)

    with c2:
        st.markdown('<div class="chart-card">', unsafe_allow_html=True)
        fpr, tpr, _ = roc_curve(y_all, oof_probs)
        roc_auc = auc(fpr, tpr)
        fig9 = go.Figure()
        fig9.add_trace(
            go.Scatter(
                x=fpr,
                y=tpr,
                mode="lines",
                name=f"ROC (AUC={roc_auc:.3f})",
                line=dict(color="#7c3aed", width=3),
            )
        )
        fig9.add_trace(
            go.Scatter(
                x=[0, 1],
                y=[0, 1],
                mode="lines",
                line=dict(color="#94a3b8", dash="dash"),
                showlegend=False,
            )
        )
        fig9.update_layout(
            **PL,
            title="ROC Curve (Out-of-Fold)",
            xaxis_title="False Positive Rate",
            yaxis_title="True Positive Rate",
            height=340,
        )
        st.plotly_chart(fig9, width="stretch")
        st.markdown("</div>", unsafe_allow_html=True)

    # ── Feature importance from dominant HGB model (bulletproof SHAP) ──
    st.markdown('<div class="chart-card">', unsafe_allow_html=True)
    try:
        sv = _safe_shap_global(global_shap_values(
            cv_metrics["fingerprint"], df[features].values, pipeline,
        ))
        mean_abs_shap = np.mean(np.abs(sv), axis=0)
        imp_df = pd.DataFrame({
            "Feature": selected_features,
            "Mean |SHAP|": mean_abs_shap,
        }).sort_values("Mean |SHAP|", ascending=True)
        fig10 = px.bar(
            imp_df,
            y="Feature",
            x="Mean |SHAP|",
            orientation="h",
            title=f"{len(selected_features)} Selected Features — Mean |SHAP| (HistGradientBoosting)",
            color="Mean |SHAP|",
            color_continuous_scale="Purples",
        )
        fig10.update_layout(**PL, coloraxis_showscale=False, height=420)
        st.plotly_chart(fig10, width="stretch")
    except Exception:
        try:
            importances = (
                hgb_model.feature_importances_
                if hasattr(hgb_model, 'feature_importances_')
                else np.zeros(len(selected_features))
            )
        except Exception:
            importances = np.zeros(len(selected_features))
        imp_df = pd.DataFrame({
            "Feature": selected_features,
            "Importance": importances,
        }).sort_values("Importance", ascending=True)
        fig10 = px.bar(
            imp_df,
            y="Feature",
            x="Importance",
            orientation="h",
            title="Feature Importances (HistGradientBoosting)",
            color="Importance",
            color_continuous_scale="Purples",
        )
        fig10.update_layout(**PL, coloraxis_showscale=False, height=420)
        st.plotly_chart(fig10, width="stretch")
    st.markdown("</div>", unsafe_allow_html=True)

    with st.expander(f"📋 {len(selected_features)} selected acoustic features"):
        selector = pipeline.hgb_pipeline_.named_steps["selector"]
        selection = pd.DataFrame({
            "Feature": features, "Mutual information": selector.scores_,
            "Selected": selector.get_support(),
        }).sort_values("Mutual information", ascending=False)
        st.dataframe(selection, width="stretch", hide_index=True)
        st.caption(
            "Final fit shown here. Each validation fold learns its own selection. "
            "Mutual information ranks relevance; correlated features can still be selected. "
            "SHAP explains the HGB component, not the complete stack."
        )

# ── TAB 4: PREDICT ───────────────────────────
with tab4:
    st.markdown(
        '<div class="sec-hdr">🧬 Interactive Prediction — Adjust Voice Features</div>',
        unsafe_allow_html=True,
    )
    st.info(
        "Use the sliders to simulate a voice recording and get a real-time "
        "model score from the grouped stack. Only the selected features are shown."
    )

    defaults = df[features].mean().to_dict()
    mins = df[features].min().to_dict()
    maxs = df[features].max().to_dict()

    with st.form("prediction_inputs", border=False):
        input_vals = defaults.copy()
        groups_ui = [
            ("🎵 Fundamental Frequency", FREQ_COLS),
            ("〰️ Jitter Features", JITTER_COLS),
            ("📶 Shimmer Features", SHIMMER_COLS),
            ("📡 Noise Ratios", RATIO_COLS),
            ("🌀 Nonlinear Dynamics", NONLINEAR_COLS),
        ]
        for grp_name, grp_cols in groups_ui:
            grp_cols = [f for f in grp_cols if f in selected_features]
            if not grp_cols:
                continue
            st.markdown(
                f'<div class="sec-hdr">{grp_name}</div>', unsafe_allow_html=True
            )
            gcols = st.columns(min(3, len(grp_cols)))
            for i, feat in enumerate(grp_cols):
                with gcols[i % len(gcols)]:
                    lo = float(mins[feat])
                    hi = float(maxs[feat])
                    dv = float(defaults[feat])
                    step = max((hi - lo) / 200, 1e-6)
                    input_vals[feat] = st.slider(
                        feat, lo, hi, dv, step=step, format="%.5f"
                    )

        submitted = st.form_submit_button(
            "🔍 Run Prediction", width="stretch", type="primary"
        )
    if submitted:
        x_in = np.array([[input_vals[f] for f in features]])
        pred, prob, x_scaled = predict_single(x_in)
        conf = round(prob[pred] * 100, 1)
        st.caption(f"Decision rule: Parkinson's score ≥ {decision_threshold:.3f}. Scores are not clinical confidence estimates.")

        if pred == 1:
            st.markdown(
                f"""
            <div class="pred-box pred-pos">
              <div class="pred-title">🔴 Model classification: Parkinson's</div>
              <div class="pred-conf">Class score: <strong>{conf}%</strong></div>
              <div class="pred-conf" style="margin-top:8px;font-size:.85rem">
                Healthy {round(prob[0] * 100, 1)}% · Parkinson's {round(prob[1] * 100, 1)}%
              </div>
            </div>""",
                unsafe_allow_html=True,
            )
        else:
            st.markdown(
                f"""
            <div class="pred-box pred-neg">
              <div class="pred-title">🟢 Model classification: Healthy</div>
              <div class="pred-conf">Class score: <strong>{conf}%</strong></div>
              <div class="pred-conf" style="margin-top:8px;font-size:.85rem">
                Healthy {round(prob[0] * 100, 1)}% · Parkinson's {round(prob[1] * 100, 1)}%
              </div>
            </div>""",
                unsafe_allow_html=True,
            )

        # ── SHAP Explanation ──
        st.markdown('<div class="chart-card">', unsafe_allow_html=True)
        st.markdown(
            '<div class="sec-hdr">🔍 SHAP Feature Contribution Analysis</div>',
            unsafe_allow_html=True,
        )
        render_shap_plot(x_scaled, selected_features)
        st.markdown("</div>", unsafe_allow_html=True)

    st.caption("Research use only. SHAP explains the HGB component; model scores are not a diagnosis.")

# ── TAB 5: DATA EXPLORER ─────────────────────
with tab5:
    st.markdown(f"**{len(dff):,}** records match current filters.")
    c1, c2 = st.columns([3, 1])
    with c1:
        search = st.text_input("🔍 Filter by patient ID", "")
    with c2:
        n_rows = st.slider("Rows to show", 1, max(2, len(dff)), min(50, len(dff)))

    disp = dff.copy()
    if search:
        disp = disp[disp["patient_id"].str.contains(search, case=False)]

    show_cols = ["name", "label"] + list(features)
    st.dataframe(
        disp[show_cols].head(n_rows).reset_index(drop=True),
        width="stretch",
        height=440,
    )

    st.markdown("**📊 Summary Statistics**")
    st.dataframe(dff[features].describe().T, width="stretch")

    st.download_button(
        "⬇️ Download Filtered Data (CSV)",
        data=dff.to_csv(index=False).encode("utf-8"),
        file_name="parkinsons_filtered.csv",
        mime="text/csv",
    )

# ══════════════════════════════════════════════
#  DOCX HELPERS
# ══════════════════════════════════════════════
def _set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _add_divider(doc, color="1e3a8a", sz="6"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), sz)
    bot.set(qn("w:space"), "1")
    bot.set(qn("w:color"), color)
    pBdr.append(bot)
    pPr.append(pBdr)


def _sec_hdr(doc, num, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"{num}.  {text}")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)


def generate_voice_docx(
    name, age, gender, healthy_prob, parkinson_prob, level, fo, hnr, advice,
    feature_names, feature_values, source_label="WAV Upload",
):
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(2)
        sec.bottom_margin = Cm(2)
        sec.left_margin = Cm(2.5)
        sec.right_margin = Cm(2.5)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    # BANNER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("NEUROCARE DIAGNOSTIC LABORATORY")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(4)
    r2 = p2.add_run("AI-Powered Parkinson's Voice Screening Report")
    r2.italic = True
    r2.font.size = Pt(11)
    r2.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    _add_divider(doc)

    # META
    mt = doc.add_table(rows=1, cols=4)
    mt.style = "Table Grid"
    meta_vals = [
        f"Report ID: PD-{random.randint(10000, 99999)}",
        f"Date: {datetime.now().strftime('%d %b %Y  %H:%M')}",
        f"Source: {source_label}",
        "Physician: — AI Screening —",
    ]
    for i, v in enumerate(meta_vals):
        mt.rows[0].cells[i].text = v
        _set_cell_bg(mt.rows[0].cells[i], "EFF6FF" if i % 2 == 0 else "FFFFFF")
        for run in mt.rows[0].cells[i].paragraphs[0].runs:
            run.font.size = Pt(9)
    doc.add_paragraph()

    # 1. PATIENT INFO
    _sec_hdr(doc, "1", "Patient Information")
    pt = doc.add_table(rows=2, cols=6)
    pt.style = "Table Grid"
    for i, h in enumerate(["Name", "Age", "Gender", "Model Class", "Fo (Hz)", "HNR"]):
        pt.rows[0].cells[i].text = h
        _set_cell_bg(pt.rows[0].cells[i], "1e3a8a")
        for run in pt.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    for i, v in enumerate(
        [name, str(age), gender, level, f"{fo:.2f}", f"{hnr:.2f}"]
    ):
        pt.rows[1].cells[i].text = v
        for run in pt.rows[1].cells[i].paragraphs[0].runs:
            run.font.size = Pt(10)
    doc.add_paragraph()

    # 2. RESULT
    _sec_hdr(doc, "2", "Diagnosis Result")
    rt = doc.add_table(rows=1, cols=3)
    rt.style = "Table Grid"
    pk_clr = (
        RGBColor(0xDC, 0x26, 0x26)
        if parkinson_prob >= decision_threshold
        else RGBColor(0x16, 0xA3, 0x4A)
    )
    hl_clr = (
        RGBColor(0x16, 0xA3, 0x4A)
        if parkinson_prob < decision_threshold
        else RGBColor(0xDC, 0x26, 0x26)
    )
    rh = {"Low Risk": "16a34a", "Moderate Risk": "d97706", "High Risk": "dc2626"}.get(
        level, "475569"
    )
    rsk_clr = RGBColor(int(rh[0:2], 16), int(rh[2:4], 16), int(rh[4:6], 16))
    for ci, (txt, clr, bg) in enumerate(
        [
            (
                f"Parkinson's\n{parkinson_prob * 100:.1f}%",
                pk_clr,
                "FEE2E2" if parkinson_prob >= decision_threshold else "DCFCE7",
            ),
            (f"Model Class\n{level}", rsk_clr, "F8FAFC"),
            (
                f"Healthy\n{healthy_prob * 100:.1f}%",
                hl_clr,
                "DCFCE7" if parkinson_prob < decision_threshold else "FEE2E2",
            ),
        ]
    ):
        cell = rt.rows[0].cells[ci]
        pp = cell.paragraphs[0]
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = pp.add_run(txt)
        rr.bold = True
        rr.font.size = Pt(15)
        rr.font.color.rgb = clr
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        pp.paragraph_format.space_before = Pt(8)
        pp.paragraph_format.space_after = Pt(8)
        _set_cell_bg(cell, bg)
    doc.add_paragraph()

    # 3. FEATURES
    _sec_hdr(doc, "3", "Complete Voice Feature Analysis")
    half = len(feature_names) // 2 + len(feature_names) % 2
    ft = doc.add_table(rows=half + 1, cols=4)
    ft.style = "Table Grid"
    for j, h in enumerate(["Feature", "Value", "Feature", "Value"]):
        ft.rows[0].cells[j].text = h
        _set_cell_bg(ft.rows[0].cells[j], "1e3a8a")
        for run in ft.rows[0].cells[j].paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    for i in range(half):
        row = ft.rows[i + 1]
        row.cells[0].text = feature_names[i]
        row.cells[1].text = f"{feature_values[i]:.5f}"
        ri = i + half
        if ri < len(feature_names):
            row.cells[2].text = feature_names[ri]
            row.cells[3].text = f"{feature_values[ri]:.5f}"
        bg = "EFF6FF" if i % 2 == 0 else "FFFFFF"
        for c in row.cells:
            _set_cell_bg(c, bg)
            for run in c.paragraphs[0].runs:
                run.font.size = Pt(9)
    doc.add_paragraph()

    # 4. RESEARCH INTERPRETATION
    _sec_hdr(doc, "4", "Research Interpretation")
    doc.add_paragraph(
        f"Model classification: {level}. The Parkinson's class score is "
        f"{parkinson_prob:.3f}, using a decision threshold of {decision_threshold:.3f}. "
        "These scores are not calibrated disease probabilities. "
        "Dataset cross-validation does not establish performance on uploaded audio. "
        "This output cannot establish or exclude a diagnosis."
    )
    _sec_hdr(doc, "5", "Feature Measurement Limits")
    doc.add_paragraph(
        "Nonlinear acoustic measurements are unavailable from this extractor and "
        "are listed as missing. Voice scoring is disabled if any are selected by the model. "
        "SHAP, when shown in the app, explains only the HGB component."
    )

    # 6. MODEL METRICS (actual CV metrics)
    _sec_hdr(doc, "6", "Model Performance Metrics")
    mm = doc.add_table(rows=2, cols=4)
    mm.style = "Table Grid"
    for j, h in enumerate(["Metric", "Value", "Metric", "Value"]):
        mm.rows[0].cells[j].text = h
        _set_cell_bg(mm.rows[0].cells[j], "1e3a8a")
        for run in mm.rows[0].cells[j].paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    doc_metrics = [
        ("Accuracy", f"{cv_metrics['accuracy']:.1f}%"),
        ("F1-Score", f"{cv_metrics['f1']:.1f}%"),
        ("Sensitivity", f"{cv_metrics['sensitivity']:.1f}%"),
        ("Specificity", f"{cv_metrics['specificity']:.1f}%"),
    ]
    for j, (m, v) in enumerate(doc_metrics):
        mm.rows[1].cells[j].text = f"{m}: {v}"
        for run in mm.rows[1].cells[j].paragraphs[0].runs:
            run.font.size = Pt(9)
    doc.add_paragraph()

    # DISCLAIMER
    _add_divider(doc, color="94a3b8", sz="4")
    disc = doc.add_paragraph()
    dr = disc.add_run(
        "\u26a0  DISCLAIMER: This report is generated by an AI screening tool for "
        "research/educational purposes only. It does not constitute a clinical diagnosis. "
        "Always consult a qualified medical professional."
    )
    dr.italic = True
    dr.font.size = Pt(8)
    dr.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ══════════════════════════════════════════════
#  REAL VOICE FEATURE EXTRACTION (parselmouth)
# ══════════════════════════════════════════════
def _safe_praat_val(val, default=None):
    """Preserve unavailable measurements as missing, never invented values."""
    if val is None:
        return float("nan")
    try:
        fval = float(val)
        if np.isnan(fval) or np.isinf(fval):
            return float("nan")
        return fval
    except (ValueError, TypeError):
        return float("nan")


def extract_voice_features(file_path):
    """Extract all 22 acoustic features from a WAV file using parselmouth (Praat).

    Computes real Jitter, Shimmer, HNR, NHR, and fundamental frequency metrics.
    Nonlinear dynamics features are missing. Prediction is blocked if they
    are selected. Unused input slots are filled only to satisfy the input schema.

    Returns:
        core_array: np.ndarray of shape (1, n_features) — raw features in training order
        fo: float — mean fundamental frequency
        hnr: float — harmonics-to-noise ratio
        all_feature_names: list[str] — all 22 feature names (for DOCX report)
        all_feature_values: list[float] — all 22 feature values (for DOCX report)
    """
    sound = parselmouth.Sound(file_path)
    pitch = sound.to_pitch()
    point_process = parselmouth.praat.call(
        [sound, pitch], "To PointProcess (cc)"
    )

    # ── Fundamental Frequency ──
    pv = pitch.selected_array["frequency"]
    pv = pv[pv != 0]
    if len(pv) == 0:
        raise ValueError("No voiced segment detected. Use a clear, sustained vowel recording.")
    else:
        fo = float(np.mean(pv))
        fhi = float(np.max(pv))
        flo = float(np.min(pv))

    # ── Jitter (real computation) ──
    jitter_local = _safe_praat_val(
        parselmouth.praat.call(point_process,
                              "Get jitter (local)", 0, 0, 0.0001, 0.02, 1.3),
        0.005,
    )
    jitter_local_abs = _safe_praat_val(
        parselmouth.praat.call(point_process,
                              "Get jitter (local, absolute)", 0, 0, 0.0001, 0.02, 1.3),
        0.00003,
    )
    jitter_rap = _safe_praat_val(
        parselmouth.praat.call(point_process,
                              "Get jitter (rap)", 0, 0, 0.0001, 0.02, 1.3),
        0.003,
    )
    jitter_ppq5 = _safe_praat_val(
        parselmouth.praat.call(point_process,
                              "Get jitter (ppq5)", 0, 0, 0.0001, 0.02, 1.3),
        0.003,
    )
    jitter_ddp = 3.0 * jitter_rap   # DDP = 3 × RAP

    # ── Shimmer (real computation) ──
    shimmer_local = _safe_praat_val(
        parselmouth.praat.call([sound, point_process],
                              "Get shimmer (local)", 0, 0, 0.0001, 0.02, 1.3, 1.6),
        0.03,
    )
    shimmer_local_dB = _safe_praat_val(
        parselmouth.praat.call([sound, point_process],
                              "Get shimmer (local_dB)", 0, 0, 0.0001, 0.02, 1.3, 1.6),
        0.30,
    )
    shimmer_apq3 = _safe_praat_val(
        parselmouth.praat.call([sound, point_process],
                              "Get shimmer (apq3)", 0, 0, 0.0001, 0.02, 1.3, 1.6),
        0.015,
    )
    shimmer_apq5 = _safe_praat_val(
        parselmouth.praat.call([sound, point_process],
                              "Get shimmer (apq5)", 0, 0, 0.0001, 0.02, 1.3, 1.6),
        0.02,
    )
    shimmer_apq11 = _safe_praat_val(
        parselmouth.praat.call([sound, point_process],
                              "Get shimmer (apq11)", 0, 0, 0.0001, 0.02, 1.3, 1.6),
        0.02,
    )
    shimmer_dda = 3.0 * shimmer_apq3   # DDA = 3 × APQ3

    # ── Harmonics ──
    harmonicity = sound.to_harmonicity()
    hnr_vals = harmonicity.values[harmonicity.values != -200]
    hnr = float(np.mean(hnr_vals)) if len(hnr_vals) > 0 else float("nan")
    nhr = 10 ** (-hnr / 10) if np.isfinite(hnr) else float("nan")

    # ── Unavailable nonlinear measurements ──
    # These require specialised algorithms (recurrence analysis, fractal scaling)
    # not available in parselmouth; do not fabricate measurements.
    rpde = dfa = spread1 = spread2 = d2 = ppe = float("nan")

    # ── All 22 features in dataset column order (for DOCX report) ──
    all_feature_names = [
        "MDVP:Fo(Hz)", "MDVP:Fhi(Hz)", "MDVP:Flo(Hz)",
        "MDVP:Jitter(%)", "MDVP:Jitter(Abs)", "MDVP:RAP", "MDVP:PPQ", "Jitter:DDP",
        "MDVP:Shimmer", "MDVP:Shimmer(dB)", "Shimmer:APQ3", "Shimmer:APQ5",
        "MDVP:APQ", "Shimmer:DDA",
        "NHR", "HNR",
        "RPDE", "DFA", "spread1", "spread2", "D2", "PPE",
    ]
    all_feature_values = [
        fo, fhi, flo,
        jitter_local, jitter_local_abs, jitter_rap, jitter_ppq5, jitter_ddp,
        shimmer_local, shimmer_local_dB, shimmer_apq3, shimmer_apq5,
        shimmer_apq11, shimmer_dda,
        nhr, hnr,
        rpde, dfa, spread1, spread2, d2, ppe,
    ]

    # ── All raw features in the trained order ──
    all_feats_dict = dict(zip(all_feature_names, all_feature_values))
    unavailable = [f for f in selected_features if not np.isfinite(all_feats_dict.get(f, np.nan))]
    if unavailable:
        raise ValueError("Selected features cannot be measured from audio: " + ", ".join(unavailable))
    core_values = [all_feats_dict[f] if np.isfinite(all_feats_dict.get(f, np.nan))
                   else float(df[f].median()) for f in features]
    core_array = np.array(core_values).reshape(1, -1)

    return core_array, fo, hnr, all_feature_names, all_feature_values


# ══════════════════════════════════════════════
#  TAB 6 — VOICE PREDICT
# ══════════════════════════════════════════════
with tab6:
    st.markdown(
        '<div class="sec-hdr">🎤 Voice-Based Prediction & Research Report</div>',
        unsafe_allow_html=True,
    )
    st.info(
        "Upload a WAV file or record your voice live. The model will predict "
        "an acoustic research score only when all selected features can be measured."
    )

    unmeasured_features = sorted(set(selected_features) & set(NONLINEAR_COLS))
    unsupported_features = sorted(set(features) - required_acoustic)
    voice_prediction_available = not unmeasured_features and not unsupported_features
    if not voice_prediction_available:
        st.warning(
            "Voice prediction is unavailable for this fitted model: the audio extractor cannot "
            "measure " + ", ".join(unmeasured_features + unsupported_features) + ". "
            "Use the Predict tab with measured features. Dataset medians must not stand in for your voice."
        )
    else:
        st.caption("Experimental audio workflow. CSV validation does not validate this audio extractor on new recordings.")

    # Patient info
    v1, v2, v3 = st.columns(3)
    with v1:
        v_name = st.text_input(
            "Patient Name", key="v_name", placeholder="Full name"
        )
    with v2:
        v_age = st.number_input("Age", 1, 120, 50, key="v_age")
    with v3:
        v_gender = st.selectbox(
            "Gender", ["Male", "Female", "Other"], key="v_gender"
        )

    st.markdown("---")

    # Two input columns
    col_upload, col_mid, col_record = st.columns([10, 1, 10])

    with col_upload:
        st.markdown('<div class="chart-card">', unsafe_allow_html=True)
        st.markdown("**📂 Upload a WAV file**")
        st.caption("Upload one pre-recorded .wav voice file")
        v_files = st.file_uploader(
            "wav files",
            type=["wav"],
            accept_multiple_files=False,
            label_visibility="collapsed",
            key="v_uploader",
        )
        v_files = [v_files] if v_files is not None else []
        if v_files:
            for f in v_files:
                st.audio(f)
        st.markdown("</div>", unsafe_allow_html=True)

    with col_mid:
        st.markdown("<br><br><br><br>", unsafe_allow_html=True)
        st.markdown(
            '<div style="text-align:center;font-size:1.2rem;font-weight:800;color:#94a3b8;">OR</div>',
            unsafe_allow_html=True,
        )

    with col_record:
        st.markdown('<div class="chart-card">', unsafe_allow_html=True)
        st.markdown("**🎤 Record Live Voice**")
        st.caption("Click the mic, speak clearly for 5–10 seconds, then stop")
        v_recorded = None
        if RECORDER_AVAILABLE:
            rb = audio_recorder(
                text="Click to record",
                recording_color="#e74c3c",
                neutral_color="#7c3aed",
                icon_name="microphone",
                icon_size="2x",
                pause_threshold=3.0,
                key="v_recorder",
            )
            if rb:
                st.audio(rb, format="audio/wav")
                v_recorded = rb
                st.success("✅ Recording captured!")
        else:
            st.warning(
                "Install `audio-recorder-streamlit` to enable live recording."
            )
            st.code("pip install audio-recorder-streamlit", language="bash")
        st.markdown("</div>", unsafe_allow_html=True)

    st.markdown("---")

    has_v_upload = bool(v_files)
    has_v_record = v_recorded is not None

    if not (has_v_upload or has_v_record):
        st.info(
            "⬆️ Upload a WAV file or record your voice above, then click Analyze."
        )

    if (has_v_upload or has_v_record) and st.button(
        "🔍 Analyze & Generate Report",
        width="stretch",
        type="primary",
        key="v_analyze",
        disabled=not voice_prediction_available,
    ):
        probs = []
        last_fo = 0.0
        last_hnr = 0.0
        last_fn = []
        last_fv = []
        last_x_scaled = None
        src_label = "WAV Upload"

        recordings = []
        if has_v_upload:
            recordings.extend((f.name, f.getvalue()) for f in v_files)
        if has_v_record:
            recordings.append(("Live microphone", v_recorded))
        # A report describes one recording, so don't mix its features with an
        # average prediction over different recordings.
        if len(recordings) != 1:
            st.error("Analyze one recording at a time so the score, explanation, and report refer to the same sample.")
            st.stop()
        src_label, audio_bytes = recordings[0]
        path = None
        try:
            with st.spinner("Analyzing voice…"):
                with tempfile.NamedTemporaryFile(delete=False, suffix=".wav") as tmp:
                    tmp.write(audio_bytes)
                    path = tmp.name
                core_feats, last_fo, last_hnr, last_fn, last_fv = extract_voice_features(path)
                pred, prob, last_x_scaled = predict_single(core_feats)
                probs.append(float(prob[1]))
        except (ValueError, RuntimeError, OSError) as exc:
            st.error(f"Could not analyze this recording: {exc}")
            st.stop()
        finally:
            if path is not None and os.path.exists(path):
                os.unlink(path)

        pk_prob = float(np.mean(probs))
        hl_prob = 1.0 - pk_prob

        # Metrics row
        m1, m2, m3 = st.columns(3)
        m1.metric("🔴 Parkinson's class score", f"{pk_prob * 100:.2f}%")
        m2.metric("🟢 Healthy class score", f"{hl_prob * 100:.2f}%")
        m3.metric("📊 AUC Score", f"{auc(*roc_curve(y_all, oof_probs)[:2]):.3f}")

        # ── SHAP Explanation ──
        st.markdown('<div class="chart-card">', unsafe_allow_html=True)
        st.markdown(
            '<div class="sec-hdr">🔍 SHAP Feature Contribution Analysis</div>',
            unsafe_allow_html=True,
        )
        if last_x_scaled is not None:
            render_shap_plot(last_x_scaled, selected_features)
        st.markdown("</div>", unsafe_allow_html=True)

        # Risk + result box
        if pk_prob < decision_threshold:
            level, advice = "Healthy model class", "Research output; not a clinical assessment"
            box_css = "pred-neg"
            icon = "🟢"
        else:
            level, advice = "Parkinson's model class", "Research output; not a clinical assessment"
            box_css = "pred-pos"
            icon = "🟣"

        st.markdown(
            f"""
        <div class="pred-box {box_css}">
          <div class="pred-title">{icon} {level}</div>
          <div class="pred-conf">{advice}</div>
          <div class="pred-conf" style="margin-top:8px;font-size:.85rem">
            Healthy {hl_prob * 100:.1f}% · Parkinson's {pk_prob * 100:.1f}%
          </div>
        </div>""",
            unsafe_allow_html=True,
        )

        # DOCX download
        st.markdown("---")
        doc_buf = generate_voice_docx(
            v_name,
            v_age,
            v_gender,
            hl_prob,
            pk_prob,
            level,
            last_fo,
            last_hnr,
            advice,
            last_fn,
            last_fv,
            source_label=src_label,
        )
        fname = f"PD_Report_{v_name.replace(' ', '_') or 'Patient'}_{datetime.now().strftime('%Y%m%d')}.docx"
        st.download_button(
            label="📄 Download Research Report (.docx)",
            data=doc_buf,
            file_name=fname,
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            width="stretch",
            key="v_download",
        )

    st.caption("Research use only. SHAP explains the HGB component; model scores are not a diagnosis.")

# ── FOOTER
st.markdown("---")
st.caption(
    "🧠 Parkinson's Voice Analytics · UCI ML Repository · "
    "Streamlit + Plotly + Stacking Ensemble + SHAP + joblib"
)
